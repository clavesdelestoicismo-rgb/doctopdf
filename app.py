from flask import Flask, request, send_file, jsonify
from docx import Document
import pdfkit
import os
import tempfile
import io

app = Flask(__name__, static_folder='.', static_url_path='')

# Configurar wkhtmltopdf
wkhtmltopdf_path = '/usr/bin/wkhtmltopdf'
config = pdfkit.configuration(wkhtmltopdf=wkhtmltopdf_path)

@app.route('/')
def index():
    """Servir la interfaz de usuario"""
    return send_file('index.html')

def docx_to_html(docx_path):
    """Convertir documento DOCX a HTML"""
    doc = Document(docx_path)
    html_parts = []
    
    # Añadir encabezado HTML
    html_parts.append('''
    <!DOCTYPE html>
    <html lang="es">
    <head>
        <meta charset="UTF-8">
        <style>
            body { font-family: Arial, sans-serif; padding: 20px; }
            h1, h2, h3, h4, h5, h6 { color: #333; }
            p { line-height: 1.6; margin: 10px 0; }
            table { border-collapse: collapse; width: 100%; margin: 20px 0; }
            th, td { border: 1px solid #ddd; padding: 8px; }
            th { background-color: #f2f2f2; }
            ul, ol { margin: 10px 0; padding-left: 30px; }
        </style>
    </head>
    <body>
    ''')
    
    # Procesar párrafos
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith('Heading'):
            level = int(paragraph.style.name.split()[-1]) if paragraph.style.name.split()[-1].isdigit() else 1
            html_parts.append(f'<h{level}>{paragraph.text}</h{level}>')
        else:
            if paragraph.text.strip():
                html_parts.append(f'<p>{paragraph.text}</p>')
    
    # Procesar tablas
    for table in doc.tables:
        html_parts.append('<table>')
        for row in table.rows:
            html_parts.append('<tr>')
            for cell in row.cells:
                tag = 'th' if row == table.rows[0] else 'td'
                html_parts.append(f'<{tag}>{cell.text}</{tag}>')
            html_parts.append('</tr>')
        html_parts.append('</table>')
    
    # Añadir pie de HTML
    html_parts.append('</body></html>')
    
    return ''.join(html_parts)

def doc_to_pdf(input_path, output_path):
    """Convertir documento DOC/DOCX a PDF"""
    file_ext = os.path.splitext(input_path)[1].lower()
    
    if file_ext == '.docx':
        # Convertir DOCX a HTML y luego a PDF
        html_content = docx_to_html(input_path)
        
        # Crear archivo temporal HTML
        with tempfile.NamedTemporaryFile(mode='w', suffix='.html', delete=False) as html_file:
            html_file.write(html_content)
            html_temp_path = html_file.name
        
        try:
            # Convertir HTML a PDF
            pdfkit.from_file(html_temp_path, output_path, configuration=config, options={
                'page-size': 'A4',
                'margin-top': '20mm',
                'margin-right': '20mm',
                'margin-bottom': '20mm',
                'margin-left': '20mm',
                'encoding': 'UTF-8',
                'no-outline': None
            })
        finally:
            # Limpiar archivo temporal HTML
            os.unlink(html_temp_path)
            
    elif file_ext == '.doc':
        # Para archivos .doc antiguos, usar LibreOffice si está disponible
        # o mostrar un mensaje de limitación
        try:
            import subprocess
            # Intentar convertir con libreoffice
            result = subprocess.run([
                'libreoffice', '--headless', '--convert-to', 'pdf',
                '--outdir', os.path.dirname(output_path), input_path
            ], capture_output=True, timeout=60)
            
            if result.returncode != 0:
                # Si falla, intentar mover el archivo convertido
                temp_pdf = os.path.join(os.path.dirname(output_path), 
                                       os.path.basename(input_path).replace('.doc', '.pdf'))
                if os.path.exists(temp_pdf):
                    os.rename(temp_pdf, output_path)
                else:
                    raise Exception('No se pudo convertir el archivo .doc')
        except Exception as e:
            raise Exception('Para archivos .doc antiguos, se requiere LibreOffice. Error: ' + str(e))
    else:
        raise ValueError('Formato de archivo no soportado. Use .doc o .docx')

@app.route('/convert', methods=['POST'])
def convert():
    """Endpoint para convertir documentos"""
    print(f"📥 Recibiendo solicitud POST - Método: {request.method}")
    
    if request.method != 'POST':
        return jsonify({'error': 'Método no permitido. Use POST'}), 405
    
    if 'document' not in request.files:
        print("❌ No se encontró el archivo en la solicitud")
        return jsonify({'error': 'No se encontró ningún archivo'}), 400
    
    file = request.files['document']
    
    if file.filename == '':
        return jsonify({'error': 'Nombre de archivo vacío'}), 400
    
    # Validar extensión
    allowed_extensions = ['.doc', '.docx']
    file_ext = os.path.splitext(file.filename)[1].lower()
    
    if file_ext not in allowed_extensions:
        return jsonify({'error': 'Formato no soportado. Suba un archivo .doc o .docx'}), 400
    
    try:
        # Crear archivos temporales
        print(f"📄 Procesando archivo: {file.filename}")
        with tempfile.NamedTemporaryFile(suffix=file_ext, delete=False) as temp_input:
            file.save(temp_input.name)
            input_path = temp_input.name
        
        output_filename = os.path.splitext(file.filename)[0] + '.pdf'
        
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as temp_output:
            output_path = temp_output.name
        
        # Convertir documento
        print(f"🔄 Convirtiendo {input_path} a {output_path}")
        doc_to_pdf(input_path, output_path)
        print(f"✅ Conversión exitosa")
        
        # Enviar archivo PDF
        return send_file(
            output_path,
            mimetype='application/pdf',
            as_attachment=True,
            download_name=output_filename
        )
    
    except Exception as e:
        print(f"❌ Error en conversión: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500
    
    finally:
        # Limpieza de archivos temporales
        try:
            if 'input_path' in locals():
                os.unlink(input_path)
            if 'output_path' in locals():
                os.unlink(output_path)
        except:
            pass

if __name__ == '__main__':
    print("🚀 Iniciando servidor de conversión DOC a PDF...")
    print("📄 Accede a http://localhost:5000 para usar la aplicación")
    app.run(host='0.0.0.0', port=5000, debug=False)
